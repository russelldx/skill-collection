"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import json
import os
import platform
import signal
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path
from xml.sax.saxutils import escape

from office.soffice import get_soffice_env

from openpyxl import load_workbook


@contextmanager
def setup_libreoffice_macro(filename):
    """Create a disposable profile; never discover, initialize or copy a user profile."""
    with tempfile.TemporaryDirectory(prefix="xlsx-recalc-") as directory:
        profile = Path(directory).resolve()
        basic = profile / "user/basic"
        standard = basic / "Standard"
        standard.mkdir(parents=True)
        marker = profile / "recalculated.ok"
        (basic / "script.xlc").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<library:libraries xmlns:library="http://openoffice.org/2000/library" xmlns:xlink="http://www.w3.org/1999/xlink">
  <library:library library:name="Standard" library:link="false" xlink:href="$(USER)/basic/Standard/script.xlb" xlink:type="simple"/>
</library:libraries>''', encoding="utf-8")
        (standard / "script.xlb").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false">
  <library:element library:name="Module1"/>
</library:library>''', encoding="utf-8")
        # Open via our own macro, not a CLI document argument: document macros
        # and external-link updates must never inherit a user's security settings.
        macro = f'''Sub RecalculateAndSave()
  On Error GoTo Failed
  Dim props(2) As New com.sun.star.beans.PropertyValue
  props(0).Name = "Hidden"
  props(0).Value = True
  props(1).Name = "MacroExecutionMode"
  props(1).Value = 0
  props(2).Name = "UpdateDocMode"
  props(2).Value = 0
  Dim doc As Object
  doc = StarDesktop.loadComponentFromURL("{Path(filename).resolve().as_uri()}", "_blank", 0, props())
  doc.calculateAll()
  doc.store()
  doc.close(True)
  Dim channel As Integer
  channel = FreeFile
  Open ConvertFromURL("{marker.as_uri()}") For Output As #channel
  Print #channel, "done"
  Close #channel
Failed:
  StarDesktop.terminate()
End Sub'''
        (standard / "Module1.xba").write_text(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<script:module xmlns:script="http://openoffice.org/2000/script" '
            'script:name="Module1" script:language="StarBasic">\n'
            + escape(macro) + '\n</script:module>', encoding="utf-8")
        yield profile, marker


def _run_libreoffice(cmd, timeout):
    """Own a POSIX process group so timeout cleanup also stops the LO child."""
    env = get_soffice_env() if platform.system() == "Linux" else os.environ.copy()
    process = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               text=True, env=env, start_new_session=True)
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except BaseException:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.communicate()
        raise
    return subprocess.CompletedProcess(cmd, process.returncode, stdout, stderr)


def recalc(filename, timeout=30):
    if not Path(filename).is_file():
        return {"error": f"File {filename} does not exist"}
    if platform.system() not in ("Linux", "Darwin"):
        return {"error": "Recalculation disabled on Windows/unsupported platforms: "
                         "reliable LibreOffice process-tree cleanup is not implemented. "
                         "No user profile or macros were accessed."}
    if timeout <= 0:
        return {"error": "timeout_seconds must be positive"}

    try:
        with setup_libreoffice_macro(filename) as (profile, marker):
            cmd = [
                "soffice", "-env:UserInstallation=" + profile.as_uri(),
                "--headless", "--norestore", "--nodefault", "--nofirststartwizard",
                "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
            ]
            result = _run_libreoffice(cmd, timeout)
            if result.returncode != 0:
                return {"error": result.stderr or f"LibreOffice exited {result.returncode}"}
            if not marker.is_file():
                return {"error": "LibreOffice did not confirm recalculation; no success assumed"}
    except subprocess.TimeoutExpired:
        return {"error": f"LibreOffice timed out after {timeout} seconds"}
    except Exception as exc:
        return {"error": f"Isolated LibreOffice recalculation failed: {exc}"}

    try:
        wb = load_workbook(filename, data_only=True)

        excel_errors = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
        ]
        error_details = {err: [] for err in excel_errors}
        total_errors = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        for err in excel_errors:
                            if err in cell.value:
                                location = f"{sheet_name}!{cell.coordinate}"
                                error_details[err].append(location)
                                total_errors += 1
                                break

        wb.close()

        result = {
            "status": "success" if total_errors == 0 else "errors_found",
            "total_errors": total_errors,
            "error_summary": {},
        }

        for err_type, locations in error_details.items():
            if locations:
                result["error_summary"][err_type] = {
                    "count": len(locations),
                    "locations": locations[:20],  
                }

        wb_formulas = load_workbook(filename, data_only=False)
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):
                        formula_count += 1
        wb_formulas.close()

        result["total_formulas"] = formula_count

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        sys.exit(1)

    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = recalc(filename, timeout)
    print(json.dumps(result, indent=2))
    sys.exit(2 if "error" in result else 1 if result.get("total_errors") else 0)


if __name__ == "__main__":
    main()
